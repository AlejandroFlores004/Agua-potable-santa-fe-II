import io
from decimal import Decimal

from django.db.models import Prefetch, Q
from django.template.loader import render_to_string
from django.utils import timezone
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from xhtml2pdf import pisa

from billing.models import Fee, Invoice, InvoiceLine
from connection.models import Line, Location
from customer.models import Customer

# Por ahora la descripción de la mecha se usa como campo de observaciones.
WITH_OBSERVATIONS = ~Q(description__isnull=True) & ~Q(description__exact="")

LINE_STATUSES = {"active": "Activas", "inactive": "Inactivas", "all": "Todas"}


def render_pdf(template_name, context):
    """Genera un PDF desde una plantilla. Devuelve los bytes o None si xhtml2pdf falla."""
    html = render_to_string(template_name, {**context, "generated_at": timezone.localtime()})
    buffer = io.BytesIO()
    if pisa.CreatePDF(html, dest=buffer, encoding="utf-8").err:
        return None
    return buffer.getvalue()


def receipt_status(invoice_line):
    # Mismo criterio que la tabla de recibos de la factura.
    if invoice_line.isPaid:
        return "Pagado"
    if invoice_line.is_carried:
        return "Trasladado"
    if invoice_line.invoice.is_closed:
        return "Vencido"
    return "Pendiente"


def _with_status(invoice_lines):
    invoice_lines = list(invoice_lines)
    for invoice_line in invoice_lines:
        invoice_line.status_label = receipt_status(invoice_line)
    return invoice_lines


def _pending_receipts():
    # Un recibo trasladado ya está sumado en el saldo anterior del que lo absorbió,
    # así que la deuda real son los impagos sin trasladar.
    return InvoiceLine.objects.filter(isPaid=False, carried_to__isnull=True)


def customers_report():
    return {
        "customers": Customer.objects.prefetch_related("lines").order_by("firstName", "lastName"),
    }


def customer_observations_report():
    observed_lines = Line.objects.filter(WITH_OBSERVATIONS).select_related("location").order_by("code")
    customers = (
        Customer.objects.filter(lines__in=observed_lines)
        .distinct()
        .prefetch_related(Prefetch("lines", queryset=observed_lines, to_attr="observed_lines"))
        .order_by("firstName", "lastName")
    )
    return {"customers": customers, "lines_count": observed_lines.count()}


def customer_statement(customer):
    lines = customer.lines.select_related("location", "fee").order_by("code")
    receipts = _with_status(
        InvoiceLine.objects.filter(line__customer=customer)
        .select_related("invoice", "line", "carried_to")
        .order_by("-invoice__start_date", "line__code")
    )
    pending = [receipt for receipt in receipts if receipt.status_label in ("Pendiente", "Vencido")]
    return {
        "customer": customer,
        "lines": lines,
        "receipts": receipts,
        "pending_total": sum((receipt.total for receipt in pending), Decimal("0")),
        "overdue_total": sum(
            (receipt.total for receipt in pending if receipt.status_label == "Vencido"), Decimal("0")
        ),
    }


def lines_report(status="active", location_id="", fee_id=""):
    lines = Line.objects.select_related("customer", "location", "fee").order_by(
        "location__name", "code"
    )
    filters = []
    if status == "active":
        lines = lines.filter(isActive=True)
    elif status == "inactive":
        lines = lines.filter(isActive=False)
    filters.append(f"Estado: {LINE_STATUSES.get(status, LINE_STATUSES['all'])}")

    if location_id == "none":
        lines = lines.filter(location__isnull=True)
        filters.append("Ubicación: sin ubicación")
    elif location_id:
        location = Location.objects.filter(pk=location_id).first()
        if location:
            lines = lines.filter(location=location)
            filters.append(f"Ubicación: {location}")

    if fee_id:
        fee = Fee.objects.filter(pk=fee_id).first()
        if fee:
            lines = lines.filter(fee=fee)
            filters.append(f"Cuota: {fee.name}")

    lines = list(lines)
    return {
        "lines": lines,
        "filters": filters,
        "monthly_total": sum((line.fee.amount for line in lines if line.isActive and line.fee), Decimal("0")),
    }


def debtors_report(only_overdue=False):
    receipts = _pending_receipts()
    if only_overdue:
        receipts = receipts.filter(invoice__status=Invoice.Status.CLOSED)
    receipts = _with_status(
        receipts.select_related("invoice", "line", "line__customer", "line__location").order_by(
            "line__customer__firstName", "line__customer__lastName", "line__code", "invoice__start_date"
        )
    )

    debtors = {}
    for receipt in receipts:
        customer = receipt.line.customer
        debtor = debtors.setdefault(customer.pk, {"customer": customer, "receipts": [], "total": Decimal("0")})
        debtor["receipts"].append(receipt)
        debtor["total"] += receipt.total

    return {
        "debtors": list(debtors.values()),
        "only_overdue": only_overdue,
        "grand_total": sum((debtor["total"] for debtor in debtors.values()), Decimal("0")),
        "receipts_count": len(receipts),
    }


MONEY_FORMAT = '"$"#,##0.00'
DATE_FORMAT = "dd/mm/yyyy"
DATETIME_FORMAT = "dd/mm/yyyy hh:mm"
HEADER_FONT = Font(bold=True, color="FFFFFF")
HEADER_FILL = PatternFill("solid", fgColor="1F4E78")


def _write_sheet(sheet, columns, rows, totals=None):
    """Llena una hoja: `columns` es [(encabezado, formato, ancho)], `rows` listas de valores."""
    sheet.append([header for header, _, _ in columns])
    for cell in sheet[1]:
        cell.font = HEADER_FONT
        cell.fill = HEADER_FILL
        cell.alignment = Alignment(vertical="center", wrap_text=True)
    for row in rows:
        sheet.append(row)
    if totals is not None:
        sheet.append(totals)
        for cell in sheet[sheet.max_row]:
            cell.font = Font(bold=True)

    for index, (_, number_format, width) in enumerate(columns, start=1):
        letter = get_column_letter(index)
        sheet.column_dimensions[letter].width = width
        if number_format:
            for (cell,) in sheet.iter_rows(min_row=2, min_col=index, max_col=index):
                cell.number_format = number_format
    sheet.freeze_panes = "A2"
    sheet.auto_filter.ref = f"A1:{get_column_letter(len(columns))}{len(rows) + 1}"


def _local(value):
    # Excel no acepta fechas con zona horaria.
    return timezone.localtime(value).replace(tzinfo=None) if value else None


def invoices_excel(invoice_id="", status=""):
    """Libro de Excel con todas las facturas (hoja Resumen) y todos sus recibos (hoja Recibos)."""
    invoices = Invoice.objects.select_related("account", "created_by").order_by("-start_date")
    filters = []
    if invoice_id:
        invoices = invoices.filter(pk=invoice_id)
    if status in Invoice.Status.values:
        invoices = invoices.filter(status=status)
        filters.append(f"Estado: {Invoice.Status(status).label}")
    invoices = list(invoices)

    receipts = _with_status(
        InvoiceLine.objects.filter(invoice__in=invoices)
        .select_related(
            "invoice", "line", "line__customer", "line__location", "fee", "late_fee", "carried_to", "created_by"
        )
        .order_by("-invoice__start_date", "line__code")
    )
    by_invoice = {}
    for receipt in receipts:
        by_invoice.setdefault(receipt.invoice_id, []).append(receipt)

    workbook = Workbook()
    summary = workbook.active
    summary.title = "Resumen"
    summary_columns = [
        ("ID", None, 6), ("Factura", None, 30), ("Cuenta", None, 22), ("Número de cuenta", None, 18),
        ("Inicio", DATE_FORMAT, 12), ("Fin", DATE_FORMAT, 12), ("Vencimiento", DATE_FORMAT, 12),
        ("Estado", None, 11), ("Recibos", None, 9), ("Pagados", None, 9), ("Pendientes", None, 10),
        ("Vencidos", None, 9), ("Trasladados", None, 11), ("Cuotas", MONEY_FORMAT, 13),
        ("Saldo anterior", MONEY_FORMAT, 14), ("Moras", MONEY_FORMAT, 12), ("Total facturado", MONEY_FORMAT, 15),
        ("Cobrado", MONEY_FORMAT, 13), ("Por cobrar", MONEY_FORMAT, 13), ("Creada", DATETIME_FORMAT, 16),
        ("Creada por", None, 14),
    ]
    summary_rows = []
    for invoice in invoices:
        lines = by_invoice.get(invoice.pk, [])
        count = lambda label: sum(1 for line in lines if line.status_label == label)
        total = lambda values: sum(values, Decimal("0"))
        summary_rows.append([
            invoice.pk, invoice.title, invoice.account.name, invoice.account.number,
            invoice.start_date, invoice.end_date, invoice.due_date, invoice.get_status_display(),
            len(lines), count("Pagado"), count("Pendiente"), count("Vencido"), count("Trasladado"),
            total(line.amount for line in lines),
            total(line.previous_balance for line in lines),
            total(line.late_fee_amount for line in lines),
            total(line.total for line in lines),
            total(line.total for line in lines if line.isPaid),
            total(line.total for line in lines if line.status_label in ("Pendiente", "Vencido")),
            _local(invoice.created_at), invoice.created_by.get_username(),
        ])
    summary_totals = ["", "Totales", "", "", None, None, None, ""] + [
        sum(row[index] for row in summary_rows) for index in range(8, 19)
    ] + [None, ""]
    _write_sheet(summary, summary_columns, summary_rows, summary_totals)

    detail = workbook.create_sheet("Recibos")
    detail_columns = [
        ("Factura", None, 30), ("Estado factura", None, 13), ("Vencimiento", DATE_FORMAT, 12),
        ("Código recibo", None, 22), ("Mecha", None, 12), ("Ubicación", None, 16), ("Cliente", None, 28),
        ("DUI", None, 12), ("Teléfono", None, 12), ("Correo", None, 24), ("Cuota", None, 16),
        ("Monto cuota", MONEY_FORMAT, 12), ("Saldo anterior", MONEY_FORMAT, 14), ("Mora", None, 14),
        ("Monto mora", MONEY_FORMAT, 12), ("Total", MONEY_FORMAT, 12), ("Pagado", None, 8),
        ("Estado recibo", None, 13), ("Trasladado a", None, 22), ("Creado", DATETIME_FORMAT, 16),
        ("Creado por", None, 14),
    ]
    detail_rows = []
    for receipt in receipts:
        customer = receipt.line.customer
        detail_rows.append([
            receipt.invoice.title, receipt.invoice.get_status_display(), receipt.invoice.due_date,
            receipt.code, receipt.line.code, str(receipt.line.location or ""),
            f"{customer.firstName} {customer.lastName}", customer.dui or "", customer.number or "",
            customer.email or "", receipt.fee.name if receipt.fee else "",
            receipt.amount, receipt.previous_balance, receipt.late_fee.name if receipt.late_fee else "",
            receipt.late_fee_amount, receipt.total, "Sí" if receipt.isPaid else "No",
            receipt.status_label, receipt.carried_to.code if receipt.carried_to else "",
            _local(receipt.created_at), receipt.created_by.get_username(),
        ])
    detail_totals = ["Totales"] + [""] * 10 + [
        sum((row[11] for row in detail_rows), Decimal("0")),
        sum((row[12] for row in detail_rows), Decimal("0")),
        "",
        sum((row[14] for row in detail_rows), Decimal("0")),
        sum((row[15] for row in detail_rows), Decimal("0")),
    ] + [""] * 5
    _write_sheet(detail, detail_columns, detail_rows, detail_totals)

    info = workbook.create_sheet("Información")
    info.append(["Reporte de facturas"])
    info["A1"].font = Font(bold=True, size=14)
    info.append(["Generado", _local(timezone.now())])
    info["B2"].number_format = DATETIME_FORMAT
    info.append(["Filtros", ", ".join(filters) or ("Una factura" if invoice_id else "Todas las facturas")])
    info.append([])
    info.append(["Notas"])
    info.append(["Total facturado incluye el saldo anterior trasladado, así que un mismo monto puede aparecer en varias facturas."])
    info.append(["Por cobrar solo suma recibos sin pagar que no se trasladaron a un recibo posterior."])
    info.column_dimensions["A"].width = 14
    info.column_dimensions["B"].width = 30

    buffer = io.BytesIO()
    workbook.save(buffer)
    return buffer.getvalue()


def report_form_options():
    return {
        "invoices": Invoice.objects.order_by("-start_date"),
        "invoice_statuses": Invoice.Status.choices,
        "customers": Customer.objects.order_by("firstName", "lastName"),
        "locations": Location.objects.all(),
        "fees": Fee.objects.order_by("name"),
        "line_statuses": LINE_STATUSES,
        "observations_count": Customer.objects.filter(lines__in=Line.objects.filter(WITH_OBSERVATIONS))
        .distinct()
        .count(),
    }
