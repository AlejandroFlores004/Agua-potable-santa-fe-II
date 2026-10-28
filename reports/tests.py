from datetime import date
from io import BytesIO

from django.contrib.auth.models import Group, User
from django.test import TestCase
from django.urls import reverse
from openpyxl import load_workbook

from accounts.permissions import Level, set_group_level
from billing.models import Account, Fee, Invoice, InvoiceLine
from connection.models import Line, Location
from customer.models import Customer

from . import services


class ReportsTestCase(TestCase):
    def setUp(self):
        self.user = User.objects.create_superuser("admin", password="pass")
        self.client.force_login(self.user)
        self.location = Location.objects.create(name="Centro")
        self.fee = Fee.objects.create(name="Normal", amount=5, created_by=self.user)
        self.ana = Customer.objects.create(firstName="Ana", lastName="Pérez", created_by=self.user)
        self.luis = Customer.objects.create(firstName="Luis", lastName="Gómez", created_by=self.user)
        self.ana_line = Line.objects.create(
            customer=self.ana, location=self.location, fee=self.fee,
            description="No se tiene contrato", created_by=self.user,
        )
        self.ana_plain_line = Line.objects.create(customer=self.ana, fee=self.fee, description="", created_by=self.user)
        self.luis_line = Line.objects.create(
            customer=self.luis, fee=self.fee, isActive=False, created_by=self.user,
        )

    def _invoice(self, start, status=Invoice.Status.DRAFT):
        account, _ = Account.objects.get_or_create(number="1", defaults={"name": "Banco", "created_by": self.user})
        return Invoice.objects.create(
            account=account, start_date=start, end_date=start, due_date=start,
            status=status, created_by=self.user,
        )


class ServicesTests(ReportsTestCase):
    def test_observations_only_lists_customers_with_described_lines(self):
        context = services.customer_observations_report()
        customers = list(context["customers"])
        self.assertEqual(customers, [self.ana])
        self.assertEqual(customers[0].observed_lines, [self.ana_line])
        self.assertEqual(context["lines_count"], 1)

    def test_lines_report_filters(self):
        active = services.lines_report("active")["lines"]
        self.assertCountEqual(active, [self.ana_line, self.ana_plain_line])
        self.assertEqual(services.lines_report("inactive")["lines"], [self.luis_line])
        self.assertEqual(services.lines_report("all", location_id=str(self.location.pk))["lines"], [self.ana_line])
        self.assertCountEqual(
            services.lines_report("all", location_id="none")["lines"], [self.ana_plain_line, self.luis_line]
        )
        self.assertEqual(services.lines_report("active")["monthly_total"], 10)

    def test_statement_and_debtors_ignore_carried_receipts(self):
        closed = self._invoice(date(2026, 1, 1), Invoice.Status.CLOSED)
        old = InvoiceLine.objects.create(invoice=closed, line=self.ana_line, created_by=self.user)
        current = self._invoice(date(2026, 2, 1), Invoice.Status.OPEN)
        new = InvoiceLine.objects.create(invoice=current, line=self.ana_line, created_by=self.user)
        old.refresh_from_db()
        self.assertEqual(old.carried_to, new)

        statement = services.customer_statement(self.ana)
        self.assertEqual([r.status_label for r in statement["receipts"]], ["Pendiente", "Trasladado"])
        self.assertEqual(statement["pending_total"], new.total)
        self.assertEqual(statement["overdue_total"], 0)

        debtors = services.debtors_report()
        self.assertEqual(len(debtors["debtors"]), 1)
        self.assertEqual(debtors["grand_total"], new.total)
        self.assertEqual(services.debtors_report(only_overdue=True)["debtors"], [])

    def test_invoices_excel(self):
        closed = self._invoice(date(2026, 1, 1), Invoice.Status.CLOSED)
        old = InvoiceLine.objects.create(invoice=closed, line=self.ana_line, created_by=self.user)
        InvoiceLine.objects.create(invoice=closed, line=self.luis_line, isPaid=True, created_by=self.user)
        current = self._invoice(date(2026, 2, 1), Invoice.Status.OPEN)
        new = InvoiceLine.objects.create(invoice=current, line=self.ana_line, created_by=self.user)

        workbook = load_workbook(BytesIO(services.invoices_excel()))
        self.assertEqual(workbook.sheetnames, ["Resumen", "Recibos", "Información"])

        summary = list(workbook["Resumen"].iter_rows(min_row=2, values_only=True))
        # Más reciente primero, más la fila de totales.
        self.assertEqual([row[0] for row in summary], [current.pk, closed.pk, None])
        self.assertEqual(summary[0][8:13], (1, 0, 1, 0, 0))
        self.assertEqual(summary[1][8:13], (2, 1, 0, 0, 1))
        self.assertEqual(summary[0][18], float(new.total))
        self.assertEqual(summary[1][17], 5)
        self.assertEqual(summary[1][18], 0)

        receipts = list(workbook["Recibos"].iter_rows(min_row=2, values_only=True))
        self.assertEqual(len(receipts), 4)
        old.refresh_from_db()
        old_row = next(row for row in receipts if row[3] == old.code)
        self.assertEqual(old_row[6], "Ana Pérez")
        self.assertEqual(old_row[17], "Trasladado")
        self.assertEqual(old_row[18], new.code)
        self.assertEqual(receipts[-1][15], float(sum(r.total for r in InvoiceLine.objects.all())))

        only_open = load_workbook(BytesIO(services.invoices_excel(status="open")))
        self.assertEqual(only_open["Recibos"].max_row, 3)  # Encabezado, un recibo y totales.
        only_closed = load_workbook(BytesIO(services.invoices_excel(invoice_id=str(closed.pk))))
        self.assertEqual(only_closed["Resumen"].cell(2, 1).value, closed.pk)


class ViewsTests(ReportsTestCase):
    def assertPdf(self, response):
        self.assertEqual(response.status_code, 200)
        self.assertEqual(response["Content-Type"], "application/pdf")
        self.assertTrue(response.content.startswith(b"%PDF"))

    def test_home(self):
        response = self.client.get(reverse("reports_home"))
        self.assertContains(response, "Clientes con observaciones")
        self.assertContains(response, "Listado de mechas")

    def test_pdfs(self):
        invoice = self._invoice(date(2026, 1, 1), Invoice.Status.OPEN)
        InvoiceLine.objects.create(invoice=invoice, line=self.ana_line, created_by=self.user)
        self.assertPdf(self.client.get(reverse("report_customers")))
        self.assertPdf(self.client.get(reverse("report_customer_observations")))
        self.assertPdf(self.client.get(reverse("report_customer_statement"), {"customer": self.ana.pk}))
        self.assertPdf(self.client.get(reverse("report_lines"), {"status": "all", "location": "none"}))
        self.assertPdf(self.client.get(reverse("report_debtors"), {"overdue": "1"}))

    def test_invoices_excel(self):
        response = self.client.get(reverse("report_invoices_excel"), {"invoice": "abc", "status": "bogus"})
        self.assertEqual(response.status_code, 200)
        self.assertIn("spreadsheetml", response["Content-Type"])
        self.assertIn("attachment", response["Content-Disposition"])
        self.assertTrue(response.content.startswith(b"PK"))
        self.assertContains(self.client.get(reverse("reports_home")), "Facturas en Excel")

    def test_statement_without_customer_redirects(self):
        response = self.client.get(reverse("report_customer_statement"))
        self.assertRedirects(response, reverse("reports_home"))
        response = self.client.get(reverse("report_customer_statement"), {"customer": 999})
        self.assertEqual(response.status_code, 404)

    def test_requires_login(self):
        self.client.logout()
        response = self.client.get(reverse("report_customers"))
        self.assertEqual(response.status_code, 302)
        self.assertTrue(response.url.startswith(reverse("login")))

    def test_requires_module_permission(self):
        group = Group.objects.create(name="Fontaneros")
        set_group_level(group, "connection", Level.READ)
        user = User.objects.create_user("fontanero", password="pass")
        user.groups.add(group)
        self.client.force_login(user)

        self.assertPdf(self.client.get(reverse("report_lines")))
        response = self.client.get(reverse("report_customers"))
        self.assertEqual(response.status_code, 403)
        self.assertContains(response, "Clientes", status_code=403)
        self.assertEqual(self.client.get(reverse("report_debtors")).status_code, 403)
        self.assertEqual(self.client.get(reverse("report_invoices_excel")).status_code, 403)

        home = self.client.get(reverse("reports_home"))
        self.assertContains(home, "Listado de mechas")
        self.assertNotContains(home, "Clientes con observaciones")
